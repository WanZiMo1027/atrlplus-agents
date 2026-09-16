#!/usr/bin/env python3
"""只读分析 QQ 群 Excel 导出，生成匿名化的风格候选证据。"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl


CHAT_SHEET = "聊天记录"
TEXT_TYPES = {"文本", "回复"}
MEDIA_PATTERN = re.compile(
    r"\[(?:图片|image|视频|video|文件|file|音频|audio):[^\]]+\]", re.IGNORECASE
)
ERROR_PATTERN = re.compile(
    r"^(?:LLM\s*响应错误|API\s*错误|Error\b|Traceback\b|Network error|请求失败)",
    re.IGNORECASE,
)
PROMPT_DUMP_PATTERN = re.compile(
    r"Output field contract|必须以规定的 JSON|可调用工具|角色说明：|系统自检脚本|"
    r"character_name.{0,40}sprite",
    re.IGNORECASE | re.DOTALL,
)
CODE_PATTERN = re.compile(
    r"(?:^|\n)\s*(?:import\s+\w+|from\s+\w+\s+import|def\s+\w+\(|class\s+\w+|"
    r"const\s+\w+\s*=|function\s+\w+\s*\()",
    re.MULTILINE,
)
HARD_ATTACK_PATTERN = re.compile(
    r"cnm|操你妈|草你妈|你妈的|我给你妈|傻逼|弱智|废物|畜生|死妈",
    re.IGNORECASE,
)
MEMORY_PATTERN = re.compile(r"记得|上次|之前|前几天|那天|以前")
PRAISE_PATTERN = re.compile(r"厉害|聪明|高性能|牛逼|好强|可爱|好看|真棒", re.IGNORECASE)
SERIOUS_PATTERN = re.compile(
    r"代码|模型|接口|提示词|skill|bug|报错|错误|api|服务器|数据库|部署|算法|"
    r"和弦|乐理|扒谱|音色|效果器|吉他|谱子|调弦|医院|科室",
    re.IGNORECASE,
)
TERMINAL_PUNCT_PATTERN = re.compile(r"[。.!！?？~～…]$")
META_AI_PHRASES = (
    "我收着点",
    "收着点",
    "收着",
    "收敛点",
    "我会克制",
    "控制攻击",
    "降低攻击",
    "不攻击你",
    "我会注意",
    "我会调整",
    "切换到友好模式",
    "收到你的要求",
)


@dataclass(frozen=True)
class Message:
    sequence: int
    timestamp: str
    sender_name: str
    sender_qq: str
    message_type: str
    content: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="分析 QQ 群 Excel 导出并生成匿名化风格证据。"
    )
    parser.add_argument("inputs", nargs="+", type=Path, help="一个或多个 QQ 群 xlsx 导出")
    parser.add_argument("--target-qq", default="3690145414", help="主要正样本账号")
    parser.add_argument(
        "--target-alias",
        default="新世界若叶睦,若叶睦,睦子,小睦",
        help="逗号分隔的目标昵称和常用称呼，用于识别直接呼叫",
    )
    parser.add_argument(
        "--exclude-qq",
        default="3784423819,3763604733",
        help="逗号分隔的其他 bot 账号，仅用于群聊响应代理统计",
    )
    parser.add_argument(
        "--from-date", default="2026-07-27", help="稳定版本起始日期，YYYY-MM-DD"
    )
    parser.add_argument("--out", required=True, type=Path, help="暂存输出目录")
    parser.add_argument(
        "--max-examples", type=int, default=120, help="候选案例最大数量"
    )
    return parser.parse_args()


def parse_timestamp(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def clean_content(content: str) -> str:
    text = MEDIA_PATTERN.sub("", content)
    text = text.replace("[回复消息]", "")
    text = re.sub(r"https?://\S+", "", text)
    return text.strip()


def style_content(content: str) -> str:
    """去掉展示层的 @ 提及，避免昵称长度污染话术长度统计。"""
    return re.sub(r"@\S+", "", clean_content(content)).strip()


def style_exclusion_reason(text: str) -> str | None:
    if ERROR_PATTERN.search(text):
        return "error_or_log"
    if len(text) > 400 and PROMPT_DUMP_PATTERN.search(text):
        return "prompt_dump"
    if len(text) > 400 and CODE_PATTERN.search(text):
        return "code_dump"
    if len(text) >= 200:
        return "over_200_chars"
    return None


def load_messages(paths: Sequence[Path]) -> tuple[list[Message], list[dict[str, object]]]:
    messages: list[Message] = []
    sources: list[dict[str, object]] = []

    # 主要流程一：逐个以只读模式打开工作簿，保留原始行序和回复类型。
    for path in paths:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if CHAT_SHEET not in workbook.sheetnames:
            raise ValueError(f"{path} 缺少工作表：{CHAT_SHEET}")
        sheet = workbook[CHAT_SHEET]
        headers = [str(cell.value or "") for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        expected = ["序号", "时间", "发送者", "发送者QQ号", "群头衔", "消息类型", "消息内容"]
        if headers[:7] != expected:
            raise ValueError(f"{path} 的聊天记录列结构与预期不一致：{headers[:7]}")

        row_count = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            sequence, timestamp, sender, qq, _title, message_type, content, *_rest = row
            if sequence is None:
                continue
            row_count += 1
            messages.append(
                Message(
                    sequence=int(sequence),
                    timestamp=str(timestamp or ""),
                    sender_name=str(sender or ""),
                    sender_qq=str(qq or ""),
                    message_type=str(message_type or ""),
                    content=str(content or ""),
                )
            )
        sources.append(
            {
                "file": path.name,
                "sheets": workbook.sheetnames,
                "chat_rows": row_count,
            }
        )
        workbook.close()

    messages.sort(key=lambda item: (item.timestamp, item.sequence))
    return messages, sources


def percentile(values: Sequence[int], fraction: float) -> int | None:
    if not values:
        return None
    ordered = sorted(values)
    index = min(len(ordered) - 1, max(0, math.ceil(fraction * len(ordered)) - 1))
    return ordered[index]


def length_summary(messages: Sequence[tuple[Message, str]]) -> dict[str, object]:
    lengths = [len(text) for _message, text in messages]
    if not lengths:
        return {"count": 0}
    return {
        "count": len(lengths),
        "median": statistics.median(lengths),
        "mean": round(statistics.mean(lengths), 2),
        "p90": percentile(lengths, 0.90),
        "p95": percentile(lengths, 0.95),
        "at_most_4": sum(value <= 4 for value in lengths),
        "at_most_12": sum(value <= 12 for value in lengths),
        "over_40": sum(value > 40 for value in lengths),
        "no_terminal_punctuation": sum(
            not TERMINAL_PUNCT_PATTERN.search(text) for _message, text in messages
        ),
        "hard_attack": sum(
            HARD_ATTACK_PATTERN.search(text) is not None for _message, text in messages
        ),
    }


def target_metrics(
    messages: Sequence[Message], target_qq: str, from_date: str
) -> tuple[dict[str, object], list[tuple[Message, str]], Counter[str]]:
    all_target = [message for message in messages if message.sender_qq == target_qq]
    text_target = [message for message in all_target if message.message_type in TEXT_TYPES]
    stable: list[tuple[Message, str]] = []
    excluded: Counter[str] = Counter()
    monthly: dict[str, list[tuple[Message, str]]] = defaultdict(list)
    daily: dict[str, list[tuple[Message, str]]] = defaultdict(list)

    # 主要流程二：只把稳定版本、可见文本且非日志/粘贴内容纳入行为基线。
    for message in text_target:
        text = style_content(message.content)
        if not text:
            excluded["empty_or_media_only"] += 1
            continue
        reason = style_exclusion_reason(text)
        if reason:
            excluded[reason] += 1
            continue
        if message.timestamp[:10] < from_date:
            excluded["before_stable_version"] += 1
            continue
        pair = (message, text)
        stable.append(pair)
        monthly[message.timestamp[:7]].append(pair)
        daily[message.timestamp[:10]].append(pair)

    metrics = {
        "all_messages": len(all_target),
        "text_or_reply_messages": len(text_target),
        "first_seen": min((item.timestamp for item in all_target), default=None),
        "last_seen": max((item.timestamp for item in all_target), default=None),
        "stable_from": from_date,
        "stable_baseline": length_summary(stable),
        "by_month": {key: length_summary(value) for key, value in sorted(monthly.items())},
        "by_day": {key: length_summary(value) for key, value in sorted(daily.items())},
        "excluded": dict(excluded),
    }
    return metrics, stable, excluded


def response_proxy(
    messages: Sequence[Message],
    target_qq: str,
    excluded_qqs: set[str],
    target_aliases: Sequence[str],
) -> dict[str, dict[str, object]]:
    counts: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    alias_pattern = re.compile("|".join(re.escape(item) for item in target_aliases))

    # 主要流程三：用时间和行距构造回复倾向代理；它不是精确的回复标签。
    for index, message in enumerate(messages):
        if (
            not message.sender_qq
            or message.sender_qq == target_qq
            or message.sender_qq in excluded_qqs
            or message.message_type not in TEXT_TYPES
        ):
            continue
        content = message.content
        direct = bool(alias_pattern.search(content)) and (
            "@" in content
            or any(content.strip().startswith(alias) for alias in target_aliases)
        )
        question = re.search(r"[?？]$", content.strip()) is not None
        category = "direct" if direct else ("question" if question else "ambient")
        responded = False
        source_time = parse_timestamp(message.timestamp)
        for later in messages[index + 1 : index + 7]:
            if (parse_timestamp(later.timestamp) - source_time).total_seconds() > 60:
                break
            if later.sender_qq == target_qq:
                responded = True
                break
        counts[category][0] += 1
        counts[category][1] += int(responded)

    return {
        category: {
            "messages": values[0],
            "target_seen_within_60s_and_6_rows": values[1],
            "rate": round(values[1] / values[0], 4) if values[0] else None,
        }
        for category, values in sorted(counts.items())
    }


def phrase_audit(
    messages: Sequence[Message], target_qq: str, phrases: Sequence[str]
) -> dict[str, dict[str, int]]:
    """分别统计全群与目标账号中的可疑元话术，保留零次结果。"""
    result: dict[str, dict[str, int]] = {}

    # 主要流程四：直接审计原始文本，避免清洗步骤把用户报告的原句改掉。
    for phrase in phrases:
        all_count = sum(phrase in message.content for message in messages)
        target_count = sum(
            phrase in message.content
            for message in messages
            if message.sender_qq == target_qq
        )
        result[phrase] = {"all_messages": all_count, "target_messages": target_count}
    return result


def stable_alias(qq: str, target_qq: str, excluded_qqs: set[str]) -> str:
    if qq == target_qq:
        return "目标"
    digest = hashlib.sha256(qq.encode("utf-8")).hexdigest()[:6].upper() if qq else "SYSTEM"
    prefix = "其他BOT" if qq in excluded_qqs else "群友"
    return f"{prefix}_{digest}"


def build_name_aliases(
    messages: Sequence[Message], target_qq: str, excluded_qqs: set[str]
) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for message in messages:
        name = message.sender_name.strip()
        if not name:
            continue
        replacement = stable_alias(message.sender_qq, target_qq, excluded_qqs)
        variants = {name, re.split(r"[（(]", name, maxsplit=1)[0].strip()}
        for variant in list(variants):
            variants.add(re.sub(r"^[一二三四五六七八九十\d]+级", "", variant).strip())
        for variant in variants:
            if variant:
                aliases[variant] = replacement
    return aliases


def anonymize_text(
    text: str, name_aliases: dict[str, str], target_aliases: Sequence[str]
) -> str:
    cleaned = MEDIA_PATTERN.sub("[媒体]", text).replace("[回复消息]", "")
    # 先处理明确 @ 提及；正文只替换较长昵称，避免把昵称“我”“你”“吉他”等普通词误伤。
    for name in sorted(name_aliases, key=len, reverse=True):
        if name:
            cleaned = cleaned.replace(f"@{name}", f"@{name_aliases[name]}")
    for name in sorted(name_aliases, key=len, reverse=True):
        if len(name.strip()) >= 3:
            cleaned = cleaned.replace(name, name_aliases[name])
    for target_alias in sorted(target_aliases, key=len, reverse=True):
        cleaned = cleaned.replace(target_alias, "目标")
    cleaned = re.sub(r"\b[1-9]\d{5,11}\b", "[账号]", cleaned)
    return cleaned.strip()


def candidate_category(message: Message, text: str, from_date: str) -> str:
    if message.timestamp[:10] < from_date or style_exclusion_reason(text):
        return "anti_ai_or_legacy"
    if HARD_ATTACK_PATTERN.search(text):
        return "hard_attack"
    if MEMORY_PATTERN.search(text):
        return "memory"
    if SERIOUS_PATTERN.search(text):
        return "serious"
    if PRAISE_PATTERN.search(text):
        return "praise"
    if message.message_type == "回复":
        return "explicit_reply"
    if len(text) <= 4:
        return "very_short"
    return "ordinary_short"


def build_candidates(
    messages: Sequence[Message],
    target_qq: str,
    excluded_qqs: set[str],
    target_aliases: Sequence[str],
    from_date: str,
    maximum: int,
) -> list[dict[str, object]]:
    name_aliases = build_name_aliases(messages, target_qq, excluded_qqs)
    quotas = {
        "anti_ai_or_legacy": 12,
        "hard_attack": 12,
        "memory": 16,
        "serious": 20,
        "praise": 16,
        "explicit_reply": 20,
        "very_short": 16,
        "ordinary_short": 20,
    }
    used: Counter[str] = Counter()
    candidates: list[dict[str, object]] = []

    # 主要流程五：按行为类型限额抽取局部上下文，并在写盘前匿名化身份。
    for index, message in enumerate(messages):
        if message.sender_qq != target_qq or message.message_type not in TEXT_TYPES:
            continue
        text = clean_content(message.content)
        measured_text = style_content(message.content)
        if not measured_text:
            continue
        category = candidate_category(message, measured_text, from_date)
        if used[category] >= quotas[category] or len(candidates) >= maximum:
            continue
        context = []
        for nearby in messages[max(0, index - 3) : min(len(messages), index + 3)]:
            nearby_text = anonymize_text(nearby.content, name_aliases, target_aliases)
            if not nearby_text:
                continue
            context.append(
                {
                    "speaker": stable_alias(nearby.sender_qq, target_qq, excluded_qqs),
                    "message_type": nearby.message_type,
                    "text": nearby_text[:500],
                }
            )
        candidates.append(
            {
                "category": category,
                "date": message.timestamp[:10],
                "target_sequence": message.sequence,
                "context": context,
            }
        )
        used[category] += 1

    return candidates


def render_markdown(metrics: dict[str, object]) -> str:
    target = metrics["target"]
    stable = target["stable_baseline"]
    proxy = metrics["response_proxy"]
    phrase_counts = metrics["meta_phrase_audit"]
    lines = [
        "# QQ 群风格分析摘要",
        "",
        f"- 目标消息：{target['all_messages']}",
        f"- 目标文本/回复：{target['text_or_reply_messages']}",
        f"- 稳定版本起点：{target['stable_from']}",
        f"- 稳定样本：{stable.get('count', 0)}",
        f"- 长度中位数：{stable.get('median')}",
        f"- 长度 P90：{stable.get('p90')}",
        f"- 4 字以内：{stable.get('at_most_4', 0)}",
        f"- 12 字以内：{stable.get('at_most_12', 0)}",
        f"- 无终止标点：{stable.get('no_terminal_punctuation', 0)}",
        f"- 强攻击性表达：{stable.get('hard_attack', 0)}",
        "",
        "## 回复倾向代理",
        "",
        "| 类型 | 消息数 | 60 秒/6 行内出现目标 | 比例 |",
        "| --- | ---: | ---: | ---: |",
    ]
    for category in ("direct", "question", "ambient"):
        item = proxy.get(category, {})
        lines.append(
            f"| {category} | {item.get('messages', 0)} | "
            f"{item.get('target_seen_within_60s_and_6_rows', 0)} | {item.get('rate')} |"
        )
    lines.extend(
        [
            "",
            "> 该响应率只是时间邻近代理，不是逐消息标注。候选案例仍需人工审查。",
            "",
        ]
    )
    lines.extend(
        [
            "## 可疑元话术审计",
            "",
            "| 短语 | 全群 | 目标账号 |",
            "| --- | ---: | ---: |",
        ]
    )
    for phrase, counts in phrase_counts.items():
        lines.append(
            f"| {phrase} | {counts['all_messages']} | {counts['target_messages']} |"
        )
    lines.append("")
    return "\n".join(lines)


def write_outputs(
    output_dir: Path, metrics: dict[str, object], candidates: Iterable[dict[str, object]]
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "metrics.json").write_text(
        json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (output_dir / "candidate_examples.jsonl").open("w", encoding="utf-8") as handle:
        for candidate in candidates:
            handle.write(json.dumps(candidate, ensure_ascii=False) + "\n")
    (output_dir / "analysis.md").write_text(render_markdown(metrics), encoding="utf-8")


def main() -> None:
    args = parse_args()
    excluded_qqs = {item.strip() for item in args.exclude_qq.split(",") if item.strip()}
    target_aliases = [item.strip() for item in args.target_alias.split(",") if item.strip()]
    if not target_aliases:
        raise ValueError("--target-alias 至少需要一个非空称呼")
    messages, sources = load_messages(args.inputs)
    target, stable, excluded = target_metrics(messages, args.target_qq, args.from_date)
    metrics = {
        "sources": sources,
        "total_loaded_messages": len(messages),
        "target": target,
        "response_proxy": response_proxy(
            messages, args.target_qq, excluded_qqs, target_aliases
        ),
        "meta_phrase_audit": phrase_audit(messages, args.target_qq, META_AI_PHRASES),
        "notes": {
            "stable_sample_count": len(stable),
            "excluded_counts": dict(excluded),
            "privacy": "Metrics and candidates omit QQ numbers and replace sender names with stable aliases.",
        },
    }
    candidates = build_candidates(
        messages,
        args.target_qq,
        excluded_qqs,
        target_aliases,
        args.from_date,
        args.max_examples,
    )
    write_outputs(args.out, metrics, candidates)
    print(
        json.dumps(
            {"output": str(args.out), "metrics": metrics["target"]}, ensure_ascii=False
        )
    )


if __name__ == "__main__":
    main()
